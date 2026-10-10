"""Generate synthetic acceptance inputs locally; does not call any API or model."""

from __future__ import annotations

import json
import shutil
from pathlib import Path
from tempfile import gettempdir
from typing import Any
from openpyxl import Workbook


def prepare(directory: Path) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    cases: list[dict[str, Any]] = []

    def text(name: str, body: str) -> None:
        (directory / name).write_text(body, encoding="utf-8")

    def case(
        name: str,
        files: list[str],
        contains: list[str],
        excludes: list[str],
        **options: Any,
    ) -> None:
        cases.append(
            {
                "name": name,
                "files": [str(directory / f) for f in files],
                "expectedStatus": options.pop("expectedStatus", "COMPLETED"),
                "contains": contains,
                "excludes": excludes,
                **options,
            }
        )

    form = """# 视频创作简报
| 填写项 | 填写内容 | 备注 / 示例 |
| --- | --- | --- |
| 品牌名称 | 林间 | 示例品牌极星 |
| 产品名称 | 林间便携折叠灯 | 示例陶瓷杯 |
| 核心卖点 | 折叠灯臂可收进底座 | 示例有机认证 |

## 实际补充说明
产品规格：电池容量2400mAh，净重318g。
产品外观：苔绿色铝合金灯臂，圆角矩形底座。
"""
    text("appendix.txt", form)
    text(
        "mixed.txt",
        form
        + "\n| 部件 | 材质 | 颜色 |\n| --- | --- | --- |\n| 灯罩 | 磨砂聚碳酸酯 | 暖白色 |\n",
    )
    case(
        "混合表格正文",
        ["mixed.txt"],
        ["2400", "318", "苔绿", "聚碳酸酯", "暖白"],
        ["极星", "有机认证", "陶瓷杯"],
    )
    intro = "品牌名称：山岚。产品名称：山岚轻行双肩包。产品介绍：适用于城市通勤的双肩包。核心卖点：独立电脑隔层可容纳14英寸笔记本电脑。\n\n"
    pad = "项目背景记录：城市通勤拍摄围绕早晨出门、公共交通和办公桌展开，保持真实自然的品牌表达。产品外观为石墨灰织物包身，画面不要出现夸张的悬浮效果。\n\n"
    text(
        "long.md",
        intro
        + pad * 85
        + "产品规格：容积18L，空包重量620g。核心卖点：背板设置透气网布。\n\n"
        + pad * 85
        + "目标受众：25至40岁的城市通勤人群。核心卖点：侧袋可放置直径7cm水杯。品牌调性：克制、轻盈。视觉统一规则：石墨灰包身保持一致。",
    )
    case("长文分片", ["long.md"], ["14", "18L", "620", "透气网布", "7cm"], ["有机认证"])
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "品牌"
    for row in [
        ["填写项", "填写内容", "备注 / 示例"],
        ["品牌名称", "听雨", "示例品牌极星"],
        ["品牌调性", "温暖、克制", "示例青春活力"],
    ]:
        sheet.append(row)
    sheet = workbook.create_sheet("产品")
    for row in [
        ["填写项", "填写内容", "备注 / 示例"],
        ["产品名称", "听雨桌面收纳盒", "示例陶瓷杯"],
        ["产品规格", "长240mm、宽160mm、高90mm", "示例容量300mL"],
        [
            "核心卖点",
            "可移动隔板调整内部空间\n透明抽屉便于查看内容物\n底部硅胶脚垫减少滑动",
            "示例有机认证",
        ],
    ]:
        sheet.append(row)
    sheet = workbook.create_sheet("外观补充")
    for row in [
        ["部件", "材质", "颜色"],
        ["外壳", "ABS", "暖灰色"],
        ["抽屉", "透明PET", "透明"],
    ]:
        sheet.append(row)
    workbook.save(directory / "multi.xlsx")
    case(
        "Excel多表换行",
        ["multi.xlsx"],
        ["240", "160", "90", "ABS", "PET", "暖灰"],
        ["极星", "300mL", "有机认证", "陶瓷杯"],
        sellingPoints=3,
    )
    lines = [
        "Brand: KITE. Product: KITE folding travel stool.",
        "Description: A portable stool for outdoor sketching.",
        "Specifications: Seat height 32 cm; product weight 850 g.",
        "Selling points: Folds flat to 4 cm; removable cotton seat cover.",
        "Appearance: Navy cotton seat; silver aluminum frame.",
        "Brand tone: Calm and practical. Avoided elements: No floating objects.",
    ]
    stream = (
        b"BT /F1 12 Tf 40 760 Td 24 TL "
        + b" ".join(("(" + line + ") Tj T*").encode() for line in lines)
        + b" ET"
    )
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    ]
    content = b"%PDF-1.4\n"
    offsets = []
    for index, obj in enumerate(objects, 1):
        offsets.append(len(content))
        content += str(index).encode() + b" 0 obj\n" + obj + b"\nendobj\n"
    xref = len(content)
    content += (
        b"xref\n0 6\n0000000000 65535 f \n"
        + b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
        + f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    (directory / "travel.pdf").write_bytes(content)
    case("PDF真实提炼", ["travel.pdf"], ["KITE", "32", "850", "4"], ["certified"])
    repo = Path(__file__).resolve().parents[3]
    shutil.copyfile(
        repo / "docs/workflows/customized/templates/视频创作简报模板.docx",
        directory / "blank.docx",
    )
    case("空白模板", ["blank.docx"], [], [], expectedStatus="FAILED")
    text(
        "valid.txt",
        "品牌名称：清禾。产品名称：清禾青梅酱。产品规格：200g/瓶。核心卖点：含青梅果肉。",
    )
    (directory / "broken.pdf").write_bytes(
        b"%PDF-1.4\nThis file has no valid page tree.\n%%EOF\n"
    )
    case("部分解析失败", ["valid.txt", "broken.pdf"], [], [], expectedStatus="PARTIAL")
    text(
        "base.txt",
        "品牌名称：清禾。产品名称：清禾青梅酱。产品规格：净含量200g/瓶，保质期12个月。核心卖点：含青梅果肉。品牌调性：清新自然。",
    )
    text(
        "extra.txt",
        "产品名称：清禾青梅酱。产品外观：透明玻璃瓶、深绿色金属盖。核心卖点：瓶盖有防拆封环。产品介绍：青梅果酱。",
    )
    text(
        "conflict.txt",
        "产品名称：清禾青梅酱。产品规格：净含量250g/瓶。目标受众：喜欢酸甜口味的家庭用户。",
    )
    case(
        "多文档补充冲突",
        ["base.txt", "extra.txt", "conflict.txt"],
        ["12", "深绿", "防拆封"],
        [],
        conflict=True,
    )
    case(
        "标准模板附录",
        ["appendix.txt"],
        ["2400", "318", "苔绿"],
        ["极星", "有机认证", "陶瓷杯"],
    )
    case(
        "空模板伴随有效资料",
        ["blank.docx", "valid.txt"],
        ["清禾", "200", "青梅果肉"],
        ["暮白", "陶瓷", "300"],
        sellingPoints=1,
    )
    (directory / "cases.json").write_text(
        json.dumps(cases, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(
        f"Prepared {len(cases)} synthetic scenarios in {directory}; no API/model calls"
    )


if __name__ == "__main__":
    prepare(Path(gettempdir()) / "customized-deep-paid-20261010")
