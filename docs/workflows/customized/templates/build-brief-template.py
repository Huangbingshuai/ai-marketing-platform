"""Build the editable Word brief and example from the same field definitions."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SECTIONS = [
    ('一 品牌基础信息', [
        ('品牌名称', '例：暮白', '暮白'),
        ('项目名称', '例：暮白晨间陶瓷杯新品短片', '暮白晨间陶瓷杯新品短片'),
        ('项目背景', '例：晨间系列新品上市，向喜欢咖啡与简洁家居的人群展示从容的晨间生活，建立自然亲切的品牌印象。', '为晨间系列新品建立简洁自然的品牌印象，展现从容的晨间生活'),
        ('品牌调性', '例：自然亲切、真实克制，呈现从容的日常生活质感。', '自然亲切、真实克制，呈现从容的日常生活质感'),
        ('视觉统一规则', '例：杯身米白色、弧形把手始终一致；品牌标识保持原样；同一人物在各镜头中保持相同外观。', '品牌名称统一为暮白，标识使用官方版本；人物与产品外观全片一致，杯身米白色和弧形把手保持不变'),
        ('画面禁用元素', '例：不添加杯盖和印花，不展示保温等未确认功能，不出现其他品牌标识。', '不添加杯盖或印花，不展示保温等未确认功能，不出现其他品牌标识'),
    ]),
    ('二 产品核心信息', [
        ('产品名称', '例：暮白晨间陶瓷杯，米白款', '暮白晨间陶瓷杯，米白款'),
        ('产品介绍', '例：日常饮用器具，陶瓷杯身，适合盛放咖啡。', '日常饮用器具，陶瓷杯身，适合盛放咖啡'),
        ('产品规格', '例：容量 300 mL，单杯装；具体尺寸待确认。', '容量 300 mL，单杯装；尺寸待确认'),
        ('产品外观', '例：米白色圆润陶瓷杯身，弧形把手，无杯盖、无印花；这些特征固定，品牌标识置于杯底，包装请 AI 设计。', '米白色圆润陶瓷杯身，弧形把手，无杯盖和印花；杯型、颜色与把手保持固定，杯底使用上传的品牌标识；包装尚未确定，请 AI 根据品牌调性设计'),
        ('核心卖点', '例（每行一条）：\n弧形把手，便于持握\n300 mL 容量，适合日常盛放咖啡\n米白色杯身，便于搭配简洁餐桌', '弧形把手，便于持握\n300 mL 容量，适合日常盛放咖啡\n米白色杯身，便于搭配简洁餐桌'),
        ('目标受众', '例：25-30 岁，性别不限，全国；喜欢咖啡与简洁家居，关注生活质感，偏好安静自然的生活内容。', '25-30 岁，性别不限，全国；喜欢简洁家居与咖啡，关注日常生活质感，偏好安静自然的生活内容'),
    ]),
]


def font(style, size, bold=False):
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")


def table(doc, rows, example, reference_heading):
    t = doc.add_table(rows=1, cols=3)
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [3.5, 7.0, 6.7]
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + edge)
        for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
            border.set(qn("w:" + key), value)
        borders.append(border)
    t._tbl.tblPr.append(borders)
    for cell, title in zip(t.rows[0].cells, ["填写项", "填写内容", reference_heading]):
        cell.text = title
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for label, options, value in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, [label, value if example else "", options]):
            cell.text = text
    for index, row in enumerate(t.rows):
        # Minimum height only: filled cells and long responses can grow naturally.
        row.height = Cm(0.67 if example else 0.78)
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for column, cell in enumerate(row.cells):
            cell.width = Cm(widths[column])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = OxmlElement("w:tcMar")
            for side in ("top", "bottom", "left", "right"):
                margin = OxmlElement("w:" + side)
                margin.set(qn("w:w"), "75")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            cell._tc.get_or_add_tcPr().append(margins)
            if index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "EDF2F7")
                cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.12
                for run in p.runs:
                    run.font.size = Pt(9.5 if column == 2 else 10)
                    run.bold = index == 0
                    if column == 2 and index:
                        run.font.color.rgb = RGBColor.from_string("596579")


def build(example=False):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin, s.bottom_margin = Cm(1.65), Cm(1.65)
    s.left_margin, s.right_margin = Cm(1.9), Cm(1.9)
    normal = doc.styles["Normal"]
    font(normal, 10)
    normal.paragraph_format.space_after = Pt(5)
    font(doc.styles["Title"], 22, True)
    doc.styles["Title"].paragraph_format.space_after = Pt(9)
    font(doc.styles["Heading 1"], 13, True)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(7)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(6)
    doc.add_paragraph("视频创作简报填写示例" if example else "视频创作简报", "Title")
    if example:
        doc.add_paragraph("以下品牌、产品与文件名均为虚构，仅用于说明如何填写。")
    else:
        doc.add_paragraph("你希望做一支什么样的视频？把想法和已有资料填写在下表中即可。")
        doc.add_paragraph("在“填写内容”一栏填写，右侧均为虚构填写示例，仅供参考。没想好的可写“待确认”或“请 AI 提方案”，不涉及的可以留空；现成的产品介绍、规格说明可作为附件上传，不用重复填写。")
    for index, (heading, rows) in enumerate(SECTIONS):
        heading_paragraph = doc.add_paragraph(heading, "Heading 1")
        if index == 1:
            heading_paragraph.paragraph_format.page_break_before = True
        table(doc, rows, example, "备注 / 示例")
        if index == 1:
            appearance_note = doc.add_paragraph("产品形象由系统根据简报生成。在产品外观中写清已确定的特征和允许 AI 设计的部分；品牌 Logo 等已有素材可在页面单独上传，无需提供产品成品图片。")
            appearance_note.paragraph_format.space_before = Pt(7)
            doc.add_paragraph("涉及成分、性能、数据或效果时，提供相关说明资料即可；尚未确认的内容请注明“待确认”。")
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = "视频创作简报填写示例" if example else "视频创作简报"
    # Remove the built-in title rule and any inherited paragraph borders.
    for element in [style.element for style in doc.styles] + [p._p for p in doc.paragraphs]:
        for border in element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    destination = ROOT / ("视频创作简报填写示例.docx" if example else "视频创作简报模板.docx")
    doc.save(destination)
    # Validate editable fields and separate brand and product sections.
    check = Document(destination)
    assert len(check.tables) == 2
    assert sum(len(t.rows) - 1 for t in check.tables) == 12
    for index, t in enumerate(check.tables):
        assert [c.text for c in t.rows[0].cells] == ["填写项", "填写内容", "备注 / 示例"]
        assert all(bool(row.cells[1].text) == example for row in t.rows[1:])
    content = "\n".join(c.text for t in check.tables for row in t.rows for c in row.cells)
    assert all(row.cells[2].text.startswith("例") for t in check.tables for row in t.rows[1:])
    assert "客户" not in content
    assert all(label not in content for label in ("想重点展示的特点", "希望给谁看", "填写提示", "已有的创意或故事", "必须包含的元素", "视频风格偏好", "文案安排", "声音形式", "配乐或音效偏好", "交付内容与数量"))
    for index, t in enumerate(check.tables):
        labels = [row.cells[0].text for row in t.rows[1:]]
        assert labels == [row[0] for row in SECTIONS[index][1]]
        assert all("或" not in label and "/" not in label for label in labels)
    assert all("参考图片与补充资料" not in p.text and "其他补充" not in p.text for p in check.paragraphs)
    print(destination)


if __name__ == "__main__":
    build()
    build(example=True)
